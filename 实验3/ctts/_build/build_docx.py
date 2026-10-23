# -*- coding: utf-8 -*-
"""把 report_content.py 中的正文写进老师模板，生成实验三报告。"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from docx import Document                                    # noqa: E402
from docx.shared import Pt                                   # noqa: E402

import report_content as RC                                  # noqa: E402
from report_content import (                                 # noqa: E402
    TEMPLATE, OUT_DOC, STUDENT, DATE_TEXT,
    style_run, add_rich, body_par, add_figure, add_data_table,
)


def clear_cell(cell):
    """清空单元格内的全部段落与嵌套表格（模板中的示例内容由此移除）。"""
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)
    for t in list(cell.tables):
        t._tbl.getparent().remove(t._tbl)


def replace_paragraph_text(par, new_text):
    """整体替换段落文字：保留首个 run 的格式，其余 run 清空。"""
    runs = par.runs
    if not runs:
        par.add_run(new_text)
        return
    runs[0].text = new_text
    for r in runs[1:]:
        r.text = ""


def fill_cover(doc):
    cover = doc.tables[0]
    mapping = {3: STUDENT["班级"], 4: STUDENT["学号"], 5: STUDENT["姓名"], 6: STUDENT["小组"]}
    for row_idx, value in mapping.items():
        cell = cover.rows[row_idx].cells[1]
        par = cell.paragraphs[0]
        if par.runs:
            replace_paragraph_text(par, value)
        else:
            r = par.add_run(value)
            style_run(r, size=12)
    for p in doc.paragraphs:
        if "2026年X月X日" in p.text:
            replace_paragraph_text(p, DATE_TEXT)
            break


ENV_LINES = [
    "硬件环境：PC 计算机（内存 8GB 及以上）。",
    "软件环境：Windows 11、JDK 17.0.20.1（Temurin）、Maven 3.9.16、IntelliJ IDEA、"
    "Spring Boot 3.3.5、MyBatis-Plus 3.5.9、MySQL 8.0.29、Apifox。",
    "网络环境：可访问 Maven 中央仓库（已配置阿里云镜像）。",
]

SUMMARY = [
    "1. 对 MyBatis-Plus 定位的理解。 实验二结束时，Service 层的数据还是内存里的一个 "
    "CopyOnWriteArrayList，进程一重启就全部还原。实验三把这条“内存数据库”换成了真实 MySQL，"
    "最直观的感受是：单表 CRUD 的 SQL 一行都没有写，插入、删除、按主键查询、分页查询"
    "全部由 BaseMapper 依据实体上的映射注解自动推导生成。MyBatis-Plus 的价值并不在于"
    "“取代 MyBatis”，而在于把“单表、条件由 Wrapper 表达”这一大类重复劳动自动化，"
    "使开发者可以把精力放在复杂查询与业务逻辑上；一旦遇到分组统计这类场景，"
    "仍然要回到 MyBatis 原生注解 SQL。",

    "2. 对实体映射注解的理解。 实验前我以为 @TableName、@TableId、@TableField 只是“写全了好”，"
    "做完才发现它们是 BaseMapper 能工作的前提：框架在运行期正是靠这些注解推导表名、列名与主键策略。"
    "本次特意把实体属性取名为 name 而列名取为 project_name，制造了一处“不一致场景”。"
    "结果是 MyBatis-Plus 自动生成 SELECT id,project_name AS name,... 这样的别名，"
    "映射照常工作——这也从反面说明，注解缺失时框架并非“什么都不做”，而是按约定猜测，"
    "猜不到就会报 Unknown column 或者静默返回 null。",

    "3. 亲手踩到并解决的第一个坑：自定义注解 SQL 不走实体映射。 "
    "ProjectMapper.selectByCategory 最初只写了 @Select，返回结果里 name 恒为 null。"
    "排查后确认：@TableField(\"project_name\") 只对 MyBatis-Plus 自动生成的 SQL 生效，"
    "手写 SQL 的结果集映射由 MyBatis 原生规则处理，开启驼峰映射后 project_name 会被映射到"
    "并不存在的属性 projectName。解决办法是用 @Results/@Result 显式声明列与属性的对应关系，"
    "或在 SQL 中直接取别名。这个坑让我真正分清了“框架生成的 SQL”与“我手写的 SQL”"
    "走的是两套映射机制。",

    "4. 第二个坑：分页插件在 3.5.9 需要单独引入依赖。 首次执行 mvn clean package 时，"
    "因为只引入了 mybatis-plus-spring-boot3-starter，配置 PaginationInnerInterceptor 时"
    "找不到 SQL 解析器而失败。查阅后确认自 MyBatis-Plus 3.5.9 起，"
    "分页插件依赖的 jsqlparser 被拆分到独立的 mybatis-plus-jsqlparser 模块，必须单独声明。"
    "补上依赖后构建通过，控制台也印证了分页确实发生在数据库侧——"
    "SELECT COUNT(*) 统计总数，SELECT ... LIMIT ?,? 取当前页，而不是把全表读进内存再截取。",

    "5. 对逻辑删除与字段自动填充的体会。 用户表上加了 @TableLogic 之后，"
    "deleteById 被自动改写成 UPDATE sys_user SET deleted=1 WHERE id=? AND deleted=0，"
    "所有查询语句也被自动追加 deleted = 0 条件：接口层面表现为“删掉了”，"
    "数据层面却仍可追溯，这正是真实业务中需要的效果。"
    "与之对照，实训项目表没有加该注解，删除就是真正的 DELETE，"
    "两者在数据库里的最终状态截然不同（有效用户数 7 而全部用户数 8）。"
    "字段自动填充则把创建/更新时间从每个 Service 方法中彻底消除，"
    "这类“每张表都有、每次都要写”的字段交给统一处理器显然更可靠。",

    "6. 不足与下一步。 本次实验的接口尚未做登录鉴权与权限控制，"
    "所有接口都可被匿名调用，密码也只用 MD5 做了演示性摘要，"
    "真实项目应改用 BCrypt 等带盐慢哈希算法；此外事务目前只在单个 Service 方法内生效，"
    "跨表的一致性场景（如“新增项目的同时写入提交记录”）还没有验证。"
    "下一步计划补上登录鉴权与 Token 校验拦截器，"
    "并为 Mapper 与 Service 层编写 JUnit 5 + MockMvc 自动化测试，"
    "把本次靠手工点击验证的用例固化成可重复执行的测试。",
]


def main():
    doc = Document(TEMPLATE)
    main_tbl = doc.tables[1]

    # 1) 封面
    fill_cover(doc)

    # 2) 实验环境
    env_cell = main_tbl.rows[4].cells[0]
    clear_cell(env_cell)
    for i, line in enumerate(ENV_LINES):
        body_par(env_cell, line, indent=False, space_before=6 if i else 0)

    # 3) 实验步骤
    step_cell = main_tbl.rows[8].cells[0]
    clear_cell(step_cell)
    for block in RC.C:
        kind = block[0]
        if kind == "h":
            body_par(step_cell, block[1], bold=True, space_before=10)
        elif kind == "p":
            body_par(step_cell, block[1])
        elif kind == "fig":
            add_figure(step_cell, block[1], block[2])
        elif kind == "tbl":
            add_data_table(step_cell, block[1], block[2], block[3],
                           block[4] if len(block) > 4 else None)

    # 4) 实验总结
    sum_cell = main_tbl.rows[10].cells[0]
    clear_cell(sum_cell)
    for i, line in enumerate(SUMMARY):
        body_par(sum_cell, line, space_before=6 if i else 0)

    # 5) 保存
    os.makedirs(os.path.dirname(OUT_DOC), exist_ok=True)
    doc.save(OUT_DOC)
    print("已生成：", OUT_DOC)
    print("正文块数：", len(RC.C), " 图片：", sum(1 for b in RC.C if b[0] == "fig"),
          " 表格：", sum(1 for b in RC.C if b[0] == "tbl"))


if __name__ == "__main__":
    main()